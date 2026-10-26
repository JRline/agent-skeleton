"""Layout of an Excel workbook, so the agent reads report-style sheets correctly the first time.

Real workbooks are rarely one clean table: title rows, multi-row headers, several tables on one sheet, totals
rows and formulas. `inspect` describes each sheet as table blocks (range, header rows, label and number columns,
total rows), the formulas behind them as patterns (e.g. "E5:E10 = B - C, row by row"), and the exact call to read
each block (`read_range` in analysis code). `read_range` turns one block into a DataFrame with its header rows
joined into column names. openpyxl only (it ships with the project).
"""
from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

MAX_SCAN_ROWS = 2000          # layout is read from the first rows; the row count comes from the sheet's size
_TOTAL = re.compile(r"\b(total|sum|subtotal|grand total|合計|小計)\b", re.I)
_REF = re.compile(r"(\$?[A-Z]{1,3})\$?(\d+)")


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _grid(ws, limit: int) -> dict[tuple[int, int], object]:
    cells = {}
    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 0, limit), values_only=True), 1):
        for c, v in enumerate(row, 1):
            if v is not None and not (isinstance(v, str) and not v.strip()):
                cells[(r, c)] = v
    return cells


def _blocks(cells: dict) -> list[tuple[int, int, int, int]]:
    """Table blocks: row bands separated by empty rows, split further where an empty column separates them."""
    rows = sorted({r for r, _ in cells})
    bands, start = [], None
    for i, r in enumerate(rows):
        if start is None:
            start = r
        if i + 1 == len(rows) or rows[i + 1] > r + 1:
            bands.append((start, r))
            start = None
    out = []
    for r1, r2 in bands:
        if len({c for (r, c) in cells if r == r1}) == 1 and r2 > r1 and \
                len({c for (r, c) in cells if r1 < r <= r2}) > 1:
            out.append((r1, min(c for (r, c) in cells if r == r1), r1, min(c for (r, c) in cells if r == r1)))
            r1 += 1                                      # a one-cell first row above a wider table: its title
        cols = sorted({c for (r, c) in cells if r1 <= r <= r2})
        group = [cols[0]]
        for c in cols[1:] + [None]:
            if c is not None and c == group[-1] + 1:
                group.append(c)
                continue
            used = [r for (r, cc) in cells if r1 <= r <= r2 and group[0] <= cc <= group[-1]]
            out.append((min(used), group[0], max(used), group[-1]))      # trimmed to the rows this group uses
            if c is not None:
                group = [c]
    return out


def _describe_block(cells, formulas, r1, c1, r2, c2) -> dict:
    """Header rows, label/number columns, total rows and formula patterns of one block."""
    def row_vals(r):
        return [cells.get((r, c)) for c in range(c1, c2 + 1)]
    header_rows = []
    for r in range(r1, r2 + 1):
        vals = [v for v in row_vals(r) if v is not None]
        nums = sum(_is_num(v) for v in vals)
        labelled = cells.get((r, c1)) is not None and not _is_num(cells.get((r, c1)))
        if vals and nums >= len(vals) / 2 and labelled and r > r1:      # first data row: a label then numbers
            break
        header_rows.append(r)
    if len(header_rows) == r2 - r1 + 1:                                 # text-only block: a note or a title
        return {"kind": "text", "range": _rng(r1, c1, r2, c2), "text": " ".join(str(v) for v in row_vals(r1) if v)}
    data = [r for r in range(r1, r2 + 1) if r not in header_rows]
    label_cols = [c for c in range(c1, c2 + 1)
                  if sum(isinstance(cells.get((r, c)), str) for r in data) > len(data) / 2]
    totals = [r for r in data if any(isinstance(cells.get((r, c)), str) and _TOTAL.search(cells[(r, c)])
                                     for c in label_cols or [c1])
              or any(str(formulas.get((r, c), "")).upper().startswith("=SUM(") for c in range(c1, c2 + 1))]
    names = []
    for c in range(c1, c2 + 1):                                         # header rows joined, merged cells filled
        parts = []
        for r in header_rows:
            v = cells.get((r, c))
            k = c
            while v is None and k > c1 and r != header_rows[-1]:        # a merged / spanning header to the left
                k -= 1
                v = cells.get((r, k))
            if v is not None and str(v) not in parts:
                parts.append(str(v))
        names.append(" ".join(parts) or get_column_letter(c))
    return {"kind": "table", "range": _rng(r1, c1, r2, c2), "header_rows": header_rows, "data_rows": len(data),
            "first_data_row": data[0] if data else None, "columns": names,
            "label_columns": [get_column_letter(c) for c in label_cols], "total_rows": totals,
            "formulas": _formula_patterns(formulas, r1, c1, r2, c2)}


def _rng(r1, c1, r2, c2) -> str:
    return f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"


def _formula_patterns(formulas: dict, r1, c1, r2, c2) -> list[str]:
    """Formulas of a block, grouped: the same formula down a column (relative to its row) is one pattern."""
    by_key: dict[tuple, list] = {}
    for (r, c), f in formulas.items():
        if r1 <= r <= r2 and c1 <= c <= c2:
            rel = _REF.sub(lambda m: m.group(1) + ("{r}" if m.group(2) == str(r) else m.group(2)), f)
            by_key.setdefault((c, rel), []).append(r)
    out = []
    for (c, rel), rows in sorted(by_key.items(), key=lambda kv: (kv[0][0], kv[1][0])):
        col = get_column_letter(c)
        where = f"{col}{rows[0]}" if len(rows) == 1 else f"{col}{rows[0]}:{col}{rows[-1]}"
        out.append(f"{where} {rel.replace('{r}', 'n') if len(rows) > 1 else rel}" + (" (row by row)" if len(rows) > 1 else ""))
    return out[:12]


def inspect(path: Path, preview_rows: int = 5) -> str:
    """A readable layout of every sheet: table blocks, how to read each, formulas, and a short preview."""
    values = load_workbook(path, read_only=True, data_only=True)
    try:
        raw = load_workbook(path, read_only=True, data_only=False)
    except Exception:
        raw = None
    lines = [f"workbook {Path(path).name}: {len(values.sheetnames)} sheet(s)"]
    for name in values.sheetnames:
        ws = values[name]
        cells = _grid(ws, MAX_SCAN_ROWS)
        formulas = {}
        if raw is not None:
            for (r, c), v in _grid(raw[name], MAX_SCAN_ROWS).items():
                if isinstance(v, str) and v.startswith("="):
                    formulas[(r, c)] = v
        size = f"{ws.max_row or 0} rows x {ws.max_column or 0} columns"
        lines.append(f'\nsheet "{name}" ({size}' + (f", layout read from the first {MAX_SCAN_ROWS} rows"
                                                    if (ws.max_row or 0) > MAX_SCAN_ROWS else "") + ")")
        if not cells:
            lines.append("  (empty)")
            continue
        n = 0
        for r1, c1, r2, c2 in _blocks(cells):
            b = _describe_block(cells, formulas, r1, c1, r2, c2)
            if b["kind"] == "text":
                lines.append(f"  {b['range']}: text: {b['text'][:160]}")
                continue
            n += 1
            hdr = (f"header rows {b['header_rows'][0]}-{b['header_rows'][-1]}" if len(b["header_rows"]) > 1
                   else f"header row {b['header_rows'][0]}" if b["header_rows"] else "no header row")
            lines.append(f"  table {n}: {b['range']} · {hdr} · {b['data_rows']} data rows"
                         + (f" · labels in {', '.join(b['label_columns'])}" if b["label_columns"] else ""))
            lines.append(f"    columns: {', '.join(b['columns'])}")
            if b["total_rows"]:
                lines.append(f"    total rows (do not add them to the rows above): {', '.join(map(str, b['total_rows']))}")
            if b["formulas"]:
                lines.append("    formulas: " + "; ".join(b["formulas"]))
                uncached = [rc for rc in formulas if r1 <= rc[0] <= r2 and c1 <= rc[1] <= c2 and rc not in cells]
                if uncached:
                    lines.append(f"    {len(uncached)} formula cells have no saved value (the file was not last saved "
                                 "by Excel): recompute them in Python from the formulas above")
            lines.append(f"    read: read_range('{{path}}', '{name}', '{b['range']}', header_rows={len(b['header_rows'])})")
            try:
                df = read_range(path, name, b["range"], len(b["header_rows"]))
                lines.append("    " + df.head(preview_rows).to_csv(index=False).strip().replace("\n", "\n    "))
            except Exception as e:                                     # the layout is still useful without it
                lines.append(f"    (preview unavailable: {e})")
    return "\n".join(lines)


def read_range(path, sheet: str, cell_range: str, header_rows: int = 1) -> pd.DataFrame:
    """One table block as a DataFrame: the header rows (joined, merged headers filled) become column names."""
    c1, r1, c2, r2 = range_boundaries(cell_range)
    ws = load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = [list(row) for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2, values_only=True)]
    head, body = rows[:header_rows], rows[header_rows:]
    names = []
    for j in range(c2 - c1 + 1):
        parts = []
        for i, hr in enumerate(head):
            v, k = hr[j], j
            while v is None and k > 0 and i < len(head) - 1:            # a spanning header to the left
                k -= 1
                v = hr[k]
            if v is not None and str(v) not in parts:
                parts.append(str(v))
        names.append(" ".join(parts) or get_column_letter(c1 + j))
    seen: dict[str, int] = {}
    for j, n in enumerate(names):                                       # duplicate names break pandas
        seen[n] = seen.get(n, 0) + 1
        if seen[n] > 1:
            names[j] = f"{n} ({seen[n]})"
    body = [r for r in body if any(v is not None for v in r)]
    return pd.DataFrame(body, columns=names)
