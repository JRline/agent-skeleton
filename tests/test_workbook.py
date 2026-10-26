"""Report-style Excel workbooks: layout (titles, header rows, several tables, totals, formulas) and read_range."""
from openpyxl import Workbook

from agent import config
from agent.pysandbox import run_code
from agent.tools import ToolContext, dispatch
from agent.workbook import inspect, read_range

from tests.helpers import TempDirs


def capital_workbook(path):
    """Synthetic: a title, a two-row header, formulas with a totals row, a second table beside it, a notes sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Capital"
    ws["A1"] = "DemoBank capital summary (JPY bn, synthetic)"
    ws.merge_cells("A1:F1")
    ws["A3"], ws["B3"], ws["D3"] = "Component", "2026Q2", "2026Q3"
    ws["B4"], ws["C4"], ws["D4"], ws["E4"] = "Amount", "Weight", "Amount", "Change"
    for i, (name, before, after) in enumerate([("CET1 before deductions", 12900, 13300), ("Deductions", -1080, -1100),
                                               ("AT1", 1520, 1550)], 5):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"D{i}"], ws[f"E{i}"] = name, before, after, f"=D{i}-B{i}"
    ws["A8"], ws["B8"], ws["D8"], ws["E8"] = "Total", "=SUM(B5:B7)", "=SUM(D5:D7)", "=D8-B8"
    ws["H3"], ws["I3"], ws["H4"], ws["I4"], ws["H5"], ws["I5"] = "Ratio", "Value %", "CET1 ratio", 12.66, "Tier 1 ratio", 14.2
    wb.create_sheet("Notes")["A1"] = "Figures are illustrative."
    wb.save(path)
    return path


class Layout(TempDirs):
    def setUp(self):
        super().setUp()
        self.path = capital_workbook(config.inputs_dir() / "capital.xlsx")

    def test_titles_tables_totals_and_formulas_are_described(self):
        out = inspect(self.path)
        self.assertIn("A1:A1: text: DemoBank capital summary", out)                  # a title, not a header
        self.assertIn("table 1: A3:E8 · header rows 3-4 · 4 data rows · labels in A", out)
        self.assertIn("columns: Component, 2026Q2 Amount, 2026Q2 Weight, 2026Q3 Amount, 2026Q3 Change", out)
        self.assertIn("total rows (do not add them to the rows above): 8", out)
        self.assertIn("E5:E8 =Dn-Bn (row by row)", out)
        self.assertIn("B8 =SUM(B5:B7)", out)
        self.assertIn("have no saved value", out)                                      # written by openpyxl
        self.assertIn("table 2: H3:I5 · header row 3 · 2 data rows", out)             # the table beside it
        self.assertIn('sheet "Notes"', out)

    def test_read_range_joins_header_rows(self):
        df = read_range(self.path, "Capital", "A3:E8", header_rows=2)
        self.assertEqual(list(df.columns), ["Component", "2026Q2 Amount", "2026Q2 Weight", "2026Q3 Amount",
                                            "2026Q3 Change"])
        self.assertEqual(df["Component"].tolist(), ["CET1 before deductions", "Deductions", "AT1", "Total"])

    def test_read_file_returns_the_layout_with_ready_read_calls(self):
        out = dispatch("read_file", {"path": "inputs/capital.xlsx"}, ToolContext("r1", config.RUNS_DIR / "r1"), "r-1")
        self.assertIn("read: read_range('inputs/capital.xlsx', 'Capital', 'A3:E8', header_rows=2)", out["result"])

    def test_analysis_code_reads_a_block_inside_the_sandbox_only(self):
        ok = run_code("df = read_range('inputs/capital.xlsx', 'Capital', 'H3:I5', header_rows=1)\nresult = df", {},
                      run_id="r1", save_as="x")
        self.assertEqual(ok["result"]["rows"], [["CET1 ratio", 12.66], ["Tier 1 ratio", 14.2]])
        bad = run_code("result = read_range('/etc/passwd', 'x', 'A1:B2')", {}, run_id="r1", save_as="y")
        self.assertFalse(bad["ok"])

    def test_dates_parse_inside_analysis_code(self):
        out = run_code("import pandas as pd\nresult = str(pd.to_datetime('20261023', format='%Y%m%d').date())", {},
                       run_id="r1", save_as="z")
        self.assertEqual(out["result"]["value"], "2026-10-23")              # pandas imports _strptime internally
